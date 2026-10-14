from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

import sys
from openpyxl import load_workbook
# Usage: python tools/make_sheet.py OUTPUT.xlsx [EXISTING.xlsx to copy rows from]
out = sys.argv[1]
source = sys.argv[2] if len(sys.argv) > 2 else None
F = "Arial"
wb = Workbook()

# ---------- Instructions ----------
ins = wb.active
ins.title = "Instructions"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 3
ins.column_dimensions["B"].width = 22
ins.column_dimensions["C"].width = 80
r = 2
ins.cell(r, 2, "Wordscape GCSE: word list").font = Font(name=F, size=16, bold=True); r += 2
ins.cell(r, 2, "Fill in the Words sheet with one row per word. Fill in the yellow columns; the grey columns check each word for you.").font = Font(name=F, size=11)
ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
ins.cell(r, 2).alignment = Alignment(wrap_text=True); ins.row_dimensions[r].height = 30; r += 2

ins.cell(r, 2, "Columns").font = Font(name=F, size=12, bold=True); r += 1
yellow = PatternFill("solid", fgColor="FFF2CC")
grey = PatternFill("solid", fgColor="EDEDED")
cols = [
    ("Subject", "Fill in. Becomes the subject folder name, for example science. Spell it the same way on every row.", yellow),
    ("Topic", "Fill in. Becomes the topic folder name, for example cells. Spell it the same way on every row.", yellow),
    ("Word", "Fill in. A single word of 3 to 7 letters, A to Z only, with no spaces or hyphens. Capitals or small letters are both fine.", yellow),
    ("Clue (optional)", "Fill in if you can. A short meaning. It is read aloud when the word is solved, and shown and read aloud by the clue button.", yellow),
    ("Say as (optional)", "Only if the voice would say the word wrongly. Write it as it should be spoken, for example B. M. P. for bmp. Leave blank otherwise.", yellow),
    ("Letters", "Worked out for you: the number of letters in the word.", grey),
    ("Check", "Worked out for you: OK, or the reason the word won't fit the game.", grey),
]
thin = Side(style="thin", color="BFBFBF")
for name, desc, fill in cols:
    a = ins.cell(r, 2, name); a.font = Font(name=F, bold=True); a.fill = fill
    b = ins.cell(r, 3, desc); b.font = Font(name=F); b.alignment = Alignment(wrap_text=True, vertical="top")
    a.alignment = Alignment(vertical="top")
    for c in (a, b): c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ins.row_dimensions[r].height = 30
    r += 1
r += 1
ins.cell(r, 2, "What the Check column means").font = Font(name=F, size=12, bold=True); r += 1
checks = [
    ("OK", "The word fits the game."),
    ("Too short", "Fewer than 3 letters."),
    ("Too long", "More than 7 letters, so it can't be spelled on the 7-letter wheel. You can leave it in; it will be set aside unless a way to handle longer words is agreed."),
    ("A to Z only", "Contains a space, hyphen, number or other character."),
    ("Duplicate", "The same word appears more than once in the same topic."),
    ("Subject or topic missing", "The word has no subject or no topic."),
]
for name, desc in checks:
    a = ins.cell(r, 2, name); a.font = Font(name=F, bold=True)
    b = ins.cell(r, 3, desc); b.font = Font(name=F); b.alignment = Alignment(wrap_text=True, vertical="top")
    a.alignment = Alignment(vertical="top")
    ins.row_dimensions[r].height = 30 if len(desc) > 80 else 16
    r += 1
r += 1
ins.cell(r, 2, "Summary").font = Font(name=F, size=12, bold=True); r += 1
summ = [("Words entered", '=COUNTA(Words!C2:C1001)')] + [(n, f'=COUNTIF(Words!G2:G1001,"{n}*")') for n, _ in checks]
for name, f in summ:
    ins.cell(r, 2, name).font = Font(name=F)
    c = ins.cell(r, 3, f); c.font = Font(name=F); c.alignment = Alignment(horizontal="left")
    r += 1
r += 1
ins.cell(r, 2, "Note").font = Font(name=F, size=12, bold=True); r += 1
n = ins.cell(r, 2, ("Rows 2 to 1001 of the Words sheet are set up for checking." if source else "The three rows on the Words sheet are examples showing the format. Replace or delete them. Rows 2 to 1001 are set up for checking."))
n.font = Font(name=F, italic=True); ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)

# ---------- Words ----------
ws = wb.create_sheet("Words")
heads = ["Subject", "Topic", "Word", "Clue (optional)", "Say as (optional)", "Letters", "Check"]
widths = [16, 18, 14, 55, 18, 9, 30]
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = ws.cell(1, i, h)
    c.font = Font(name=F, bold=True)
    c.fill = yellow if i <= 5 else grey
    c.border = Border(bottom=Side(style="thin", color="808080"))
    ws.column_dimensions[c.column_letter].width = w
examples = [
    ("science", "cells", "cell", "The basic unit of all living things", None),
    ("science", "cells", "nucleus", "Controls the cell and holds its DNA", None),
    ("computer", "Digital Media", "bmp", "An uncompressed image file format", "B. M. P."),
]
N = 1001
if source:
    sw = load_workbook(source)["Words"]
    hdr = [str(c.value or "").strip().lower() for c in sw[1]]
    idx = {k: hdr.index(k) for k in ("subject", "topic", "word") if k in hdr}
    clue_i = next((i for i, h in enumerate(hdr) if h.startswith("clue")), None)
    say_i = next((i for i, h in enumerate(hdr) if h.startswith("say as")), None)
    examples = []
    for r in sw.iter_rows(min_row=2, values_only=True):
        if r[idx["word"]] in (None, ""):
            continue
        examples.append((r[idx["subject"]], r[idx["topic"]], r[idx["word"]],
                         r[clue_i] if clue_i is not None else None,
                         r[say_i] if say_i is not None else None))
for row in range(2, N + 1):
    if row - 2 < len(examples):
        for col, v in enumerate(examples[row - 2], 1):
            ws.cell(row, col, v)
    for col in range(1, 6):
        ws.cell(row, col).font = Font(name=F)
    ws.cell(row, 4).alignment = Alignment(wrap_text=False)
    ws.cell(row, 6, f'=IF(TRIM(C{row})="","",LEN(TRIM(C{row})))').font = Font(name=F, color="595959")
    w = f'TRIM(C{row})'
    ws.cell(row, 7, (
        f'=IF({w}="","",'
        f'IF(OR(TRIM(A{row})="",TRIM(B{row})=""),"Subject or topic missing",'
        f'IF(SUMPRODUCT(ISERROR(FIND(MID(UPPER({w}),ROW($1:$30),1),"ABCDEFGHIJKLMNOPQRSTUVWXYZ"))*(ROW($1:$30)<=LEN({w})))>0,"A to Z only",'
        f'IF(LEN({w})<3,"Too short (min 3 letters)",'
        f'IF(LEN({w})>7,"Too long (max 7 letters)",'
        f'IF(COUNTIFS($A$2:$A${N},A{row},$B$2:$B${N},B{row},$C$2:$C${N},C{row})>1,"Duplicate in this topic",'
        f'"OK"))))))'
    )).font = Font(name=F, color="595959")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{N}"
rng = f"G2:G{N}"
ws.conditional_formatting.add(rng, FormulaRule(formula=['$G2="OK"'], font=Font(color="006100"), fill=PatternFill("solid", fgColor="E2EFDA")))
ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($G2<>"",$G2<>"OK")'], font=Font(color="9C0006"), fill=PatternFill("solid", fgColor="FCE4E4")))
wb.active = 1
wb.save(out)
print(out)
