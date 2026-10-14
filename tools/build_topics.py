"""Build the topic word files from the word list spreadsheet.

Usage:
    python tools/build_topics.py subjects/wordscape-gcse-words.xlsx

Reads the Words sheet (Subject, Topic, Word, Clue, Say as) and writes one file per topic:
    subjects/<subject>/<topic>/words.csv   (columns: word, clue, say_as)

Folder names are the subject and topic in lower case, with spaces replaced by hyphens
("Digital Media" becomes "digital-media"). Words are stored in lower case.

Rows that fail the spreadsheet's rules (3 to 7 letters, A to Z only, subject and topic
filled in, no duplicate within a topic) are skipped and listed on screen.
Each topic's words.csv is rewritten, and subjects/topics.json lists every subject and topic. Word files for topics no longer in the spreadsheet are
removed; anything else in a topic folder (such as its own backgrounds) is left alone.
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SUBJECTS = ROOT / "subjects"
WORD = re.compile(r"^[a-z]{3,7}$")


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.strip().lower()).strip("-")


def main(path):
    ws = load_workbook(path, read_only=True)["Words"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(v or "").strip().lower() for v in next(rows)]
    col = {k: next(i for i, h in enumerate(hdr) if h.startswith(k)) for k in ("subject", "topic", "word", "clue")}
    say_i = next((i for i, h in enumerate(hdr) if h.startswith("say as")), None)
    topics, skipped, names = {}, [], {}
    for n, row in enumerate(rows, start=2):
        cell = lambda i: "" if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()
        subject, topic, word, clue = (cell(col[k]) for k in ("subject", "topic", "word", "clue"))
        say = cell(say_i)
        if not word:
            continue
        w = word.lower()
        if not subject or not topic:
            skipped.append((n, word, "subject or topic missing"))
        elif not WORD.match(w):
            skipped.append((n, word, "must be 3 to 7 letters, A to Z only"))
        elif w in topics.setdefault((slug(subject), slug(topic)), {}):
            skipped.append((n, word, "duplicate in this topic"))
        else:
            topics[(slug(subject), slug(topic))][w] = (clue, say)
            names.setdefault(slug(subject), subject)
            names.setdefault((slug(subject), slug(topic)), topic)

    # Remove word files for topics no longer in the spreadsheet. Other files in topic
    # folders (such as a topic's own backgrounds) are left alone.
    wanted = {SUBJECTS / s / tp / "words.csv" for (s, tp), w in topics.items() if w}
    for old in SUBJECTS.glob("*/*/words.csv"):
        if old not in wanted:
            old.unlink()
            print(f"Removed {old.relative_to(ROOT)} (topic no longer in the spreadsheet)")
            for d in (old.parent, old.parent.parent):
                if d.is_dir() and not any(d.iterdir()):
                    d.rmdir()
    for (subject, topic), words in sorted(topics.items()):
        if not words:
            continue
        folder = SUBJECTS / subject / topic
        folder.mkdir(parents=True, exist_ok=True)
        with (folder / "words.csv").open("w", newline="", encoding="utf-8") as fh:
            out = csv.writer(fh)
            out.writerow(["word", "clue", "say_as"])
            for w, (c, s) in words.items():
                out.writerow([w, c, s])
        print(f"{subject}/{topic}: {len(words)} words")
    # The list of subjects and topics the game reads (a web page cannot look inside folders).
    index = {"subjects": []}
    for s in sorted({s for (s, _), w in topics.items() if w}):
        entry = {"key": s, "name": names[s], "topics": []}
        for (s2, tp), words in sorted(topics.items()):
            if s2 == s and words:
                entry["topics"].append({"key": tp, "name": names[(s2, tp)], "words": len(words),
                                        "file": f"subjects/{s}/{tp}/words.csv"})
        index["subjects"].append(entry)
    (SUBJECTS / "topics.json").write_text(json.dumps(index, indent=2) + "\n")
    print("Written subjects/topics.json")
    for n, word, why in skipped:
        print(f"Skipped row {n} ({word}): {why}")


if __name__ == "__main__":
    main(sys.argv[1])
